import re
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def parse_bib(path: Path):
    text = path.read_text(encoding="utf-8")
    starts = list(re.finditer(r"(?m)^@(\w+)\s*\{\s*([^,]+),", text))
    records = {}
    for idx, match in enumerate(starts):
        end = starts[idx + 1].start() if idx + 1 < len(starts) else len(text)
        body = text[match.end() : end]
        fields = {}
        for field in ("title", "doi", "journal", "year", "author", "abstract", "url"):
            field_match = re.search(
                rf"(?ims)^\s*{field}\s*=\s*\{{(.*?)\}}\s*,?\s*$", body
            )
            if field_match:
                fields[field] = re.sub(r"\s+", " ", field_match.group(1)).strip()
        records[match.group(2).strip()] = fields
    return records


def cited_keys(path: Path):
    text = path.read_text(encoding="utf-8")
    found = set()
    for match in re.finditer(r"\\cite[pt]?\{([^}]*)\}", text):
        found.update(key.strip() for key in match.group(1).split(",") if key.strip())
    return found


def shortlist_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["Screening register"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(header, values)))
    return rows


def main():
    records = parse_bib(ROOT / "FireScape_master_v6.bib")
    cited = cited_keys(ROOT / "review_paper_v5.tex")
    rows = shortlist_rows(ROOT / "FireScape_Alps_screening_register.xlsx")
    doi_to_key = {
        value.get("doi", "").lower().strip(): key
        for key, value in records.items()
        if value.get("doi")
    }
    title_to_key = {
        re.sub(r"[^a-z0-9]+", " ", value.get("title", "").lower()).strip(): key
        for key, value in records.items()
        if value.get("title")
    }
    candidates = []
    for row in rows:
        if row.get("Decision") != "INCLUDE":
            continue
        key = doi_to_key.get(str(row.get("DOI") or "").lower().strip())
        if not key:
            normalized_title = re.sub(
                r"[^a-z0-9]+", " ", str(row.get("Title") or "").lower()
            ).strip()
            key = title_to_key.get(normalized_title)
        if key and key not in cited:
            record = records[key]
            candidates.append(
                (
                    -(row.get("Seed citations") or 0),
                    str(row.get("Tier") or ""),
                    key,
                    record.get("year", row.get("Year")),
                    record.get("title", row.get("Title")),
                    record.get("doi", row.get("DOI")),
                    row.get("Suggested section(s)"),
                )
            )
    candidates.sort()
    limit = int(sys.argv[1]) if len(sys.argv) > 1 else 150
    for item in candidates[:limit]:
        citations = -item[0]
        print("\t".join(str(value or "") for value in (citations,) + item[1:]))


if __name__ == "__main__":
    main()
